"""离屏渲染真实 GUI 分析结果，便于无桌面环境做布局检查。"""
import argparse
import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from PySide6.QtCore import QTimer
from PySide6.QtGui import QFontDatabase, QFont
from PySide6.QtWidgets import QApplication
from gui.main_window import MainWindow

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('video', nargs='?', default='samples/final_validation.mp4')
parser.add_argument('--database', default='data/gui-validation.sqlite3')
args = parser.parse_args()
app = QApplication([])
# 离屏平台不会自动发现 Windows 字体，显式加载仅用于渲染验收。
font = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts/msyh.ttc'
if font.exists():
    font_id = QFontDatabase.addApplicationFont(str(font))
    families = QFontDatabase.applicationFontFamilies(font_id)
    if families:
        app.setFont(QFont(families[0], 10))
app.setStyle('Fusion')
window = MainWindow(args.video, database_path=args.database)
window.setStyleSheet(window.styleSheet().replace('Microsoft YaHei UI', 'Microsoft YaHei'))
window.show()
projects = window.store.list_projects()
project_id = next((p['id'] for p in projects if p['name'] == '会议室投屏验收'), None)
if project_id is None:
    project_id = window.store.create_project('会议室投屏验收')
window.refresh_projects(project_id)
window.task_title_edit.setText('1080p 黑屏与冻结回归')
window.start_analysis()
exit_code = 1
stage = 'analysis'


def finish():
    global exit_code, stage
    if stage == 'analysis':
        if window.worker is not None:
            return
        if window.result is None or window.table.rowCount() != 4:
            print(window.status_label.text())
            timer.stop()
            window.close()
            return
        target = Path('docs/images/v06-analysis.png')
        target.parent.mkdir(parents=True, exist_ok=True)
        assert window.grab().save(str(target))
        print(window.result.statistics)
        window.table.selectRow(1)
        window.open_review()
        stage = 'review'
    elif stage == 'review':
        dialog = window.review_dialog
        if dialog.busy:
            return
        if dialog.current_frame is None:
            print(dialog.status.text())
        else:
            assert dialog.current_frame.index == round(25 * dialog.fps)
            assert dialog.grab().save('docs/images/v06-review.png')
            Path('screenshots').mkdir(exist_ok=True)
            dialog.save_image('screenshots/v06-freeze-25s.png')
            print(f'Frame {dialog.current_frame.index}, image {dialog.current_image.width()}x{dialog.current_image.height()}')
            exit_code = 0
        stage = 'history'
        dialog.close()
    elif stage == 'history' and window.review_dialog is None:
        from storage.analysis_store import AnalysisStore
        from gui.history_dialog import HistoryDialog
        store = AnalysisStore(args.database)
        task = store.load_task(window.task_id)
        assert task['result'].statistics == window.result.statistics
        assert task['project_id'] == project_id
        assert task['title'] == '1080p 黑屏与冻结回归'
        assert len(store.list_screenshots(window.task_id)) == 1
        window.load_history(window.task_id)
        history = HistoryDialog(store, window)
        history.project_filter.setCurrentIndex(history.project_filter.findData(project_id))
        history.show()
        app.processEvents()
        assert history.grab().save('docs/images/v06-history.png')
        print('Reopened task and screenshot association verified')
        history.close()
        history.deleteLater()
        window.export_report_to('reports/v06-validation.xlsx', overwrite=True)
        stage = 'report'
    elif stage == 'report' and window.worker is None:
        from openpyxl import load_workbook
        if '报告已导出' not in window.status_label.text():
            print(window.status_label.text())
            exit_code = 1
        else:
            book = load_workbook('reports/v06-validation.xlsx')
            summary = dict(book['分析概况'].iter_rows(min_row=2, values_only=True))
            assert summary['异常总数'] == 4
            assert summary['关联截图数'] == 1
            assert book['异常明细'].max_row == 5
            assert book['截图关联']['E2'].value == '可用'
            assert book['截图关联']['F2'].hyperlink is not None
            book.close()
            assert window.grab().save('docs/images/v06-export.png')
            print('XLSX: 4 issues, 1 verified screenshot, saved task metadata confirmed')
        timer.stop()
        window.close()
        app.quit()


timer = QTimer()
timer.timeout.connect(finish)
timer.start(100)
QTimer.singleShot(60000, window.close)
app.exec()
sys.exit(exit_code)
