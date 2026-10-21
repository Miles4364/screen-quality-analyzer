"""Build a traceable source listing without changing the application source."""
from __future__ import annotations

import argparse
import ast
import hashlib
import io
import json
import os
from pathlib import Path
import tokenize
from typing import Callable
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo


ROOT = Path(__file__).resolve().parents[1]
ENTRY_FILES = ("gui_main.py", "main.py", "launcher.py", "black_test.py")
SOURCE_DIRS = ("core", "gui", "storage", "reporting", "utils", "exceptions")
RUNTIME_TOOLS = ("tools/capture_materials.py", "tools/packaged_smoke.py")
PROJECT_FILES = (
    ".gitignore", ".gitattributes", "README.md", "LICENSE", "LICENSE.txt",
    "ScreenQualityAnalyzer.spec", "requirements.txt", "requirements-gui.txt",
    "requirements-build.txt", "requirements-docs.txt", "pyproject.toml",
    "pytest.ini", "setup.cfg", "tox.ini", "samples/README.md",
)
FONT_SIZE = 10.5
ROW_HEIGHT = 12.0
ROWS_PER_PAGE = 57
LINE_NUMBER_WIDTH = 32.0
CODE_WIDTH = (210 - 18 - 18) / 25.4 * 72 - LINE_NUMBER_WIDTH - 20


def source_files(root: Path) -> list[Path]:
    files = [root / filename for filename in ENTRY_FILES if (root / filename).is_file()]
    for dirname in SOURCE_DIRS:
        files.extend(sorted((root / dirname).rglob("*.py")))
    files.extend(root / filename for filename in RUNTIME_TOOLS if (root / filename).is_file())
    return files


def _character_column(line: str, byte_column: int) -> int:
    return len(line.encode("utf-8")[:byte_column].decode("utf-8"))


def project_source(text: str, filename: str = "<source>") -> dict:
    """Remove actual comment tokens and standalone docstrings, never string content."""
    lines = text.splitlines()
    tree = ast.parse(text, filename=filename)
    tokens = list(tokenize.generate_tokens(io.StringIO(text).readline))
    removals: dict[int, list[tuple[int, int]]] = {}
    protected_blank_rows: set[int] = set()
    docstrings: set[tuple[int, int]] = set()
    preserved_inline_docstrings = 0
    for node in ast.walk(tree):
        if not isinstance(node, (ast.Module, ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef)):
            continue
        if not node.body:
            continue
        first = node.body[0]
        if not (isinstance(first, ast.Expr) and isinstance(first.value, ast.Constant)
                and isinstance(first.value.value, str)):
            continue
        first_column = _character_column(lines[first.lineno - 1], first.col_offset)
        last_column = _character_column(lines[first.end_lineno - 1], first.end_col_offset)
        prefix = lines[first.lineno - 1][:first_column]
        suffix = lines[first.end_lineno - 1][last_column:]
        if prefix.strip() or (suffix.strip() and not suffix.lstrip().startswith("#")):
            preserved_inline_docstrings += 1
            continue
        docstrings.add((first.lineno, first.end_lineno))
        for row in range(first.lineno, first.end_lineno + 1):
            removals.setdefault(row, []).append((0, len(lines[row - 1])))
    for token in tokens:
        if token.type == tokenize.COMMENT:
            removals.setdefault(token.start[0], []).append((token.start[1], token.end[1]))
        elif token.type == tokenize.STRING and token.end[0] > token.start[0]:
            if not any(start <= token.start[0] and token.end[0] <= end for start, end in docstrings):
                protected_blank_rows.update(range(token.start[0], token.end[0] + 1))
    output = []
    for number, original in enumerate(lines, 1):
        projected = original
        for start, end in sorted(removals.get(number, []), reverse=True):
            projected = projected[:start] + projected[end:]
        projected = projected.rstrip()
        if projected.strip() or number in protected_blank_rows:
            output.append({"source_line": number, "text": projected.expandtabs(4)})
    statements = sum(
        1 for token in tokens if token.type == tokenize.NEWLINE
        and not any(start <= token.start[0] <= end for start, end in docstrings)
    )
    return {
        "lines": output,
        "physical_lines": len(lines),
        "effective_code_lines": sum(bool(line["text"].strip()) for line in output),
        "logical_statement_lines": statements,
        "preserved_inline_docstrings": preserved_inline_docstrings,
        "preserved_literal_blank_lines": sum(not line["text"].strip() for line in output),
    }


def font_measurement(latin_path: Path | None = None, chinese_path: Path | None = None) -> Callable[[str], float]:
    from PIL import ImageFont

    fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    latin = ImageFont.truetype(str(latin_path or fonts / "ARIALN.TTF"), 84)
    chinese = ImageFont.truetype(str(chinese_path or fonts / "simsun.ttc"), 84)

    def width(text: str) -> float:
        return sum((latin if ord(char) < 256 else chinese).getlength(char) / 8 for char in text)

    return width


def wrap_source_line(text: str, measure: Callable[[str], float], width: float = CODE_WIDTH) -> list[str]:
    """Soft-wrap the displayed text only; concatenating rows restores the source line."""
    if not text:
        return [""]
    rows, current = [], ""
    current_width = 0.0
    for char in text:
        char_width = measure(char)
        if current and current_width + char_width > width:
            rows.append(current)
            current, current_width = "", 0.0
        current += char
        current_width += char_width
    if current:
        rows.append(current)
    return rows


def build_plan(root: Path, measure: Callable[[str], float], rows_per_page: int = ROWS_PER_PAGE) -> dict:
    if rows_per_page < 50 or rows_per_page > 60:
        raise ValueError("rows_per_page must be between 50 and 60")
    files, listing, pages = [], [], []
    total_effective = total_physical = total_statements = 0
    for path in source_files(root):
        data = path.read_bytes()
        text = data.decode("utf-8-sig")
        relative = path.relative_to(root).as_posix()
        projection = project_source(text, relative)
        start = len(listing) + 1
        for item in projection["lines"]:
            listing.append({
                "number": len(listing) + 1, "file": relative, **item,
                "rows": wrap_source_line(item["text"], measure),
            })
        files.append({
            "path": relative, "bytes": len(data),
            "md5": hashlib.md5(data).hexdigest(),
            "sha256": hashlib.sha256(data).hexdigest(),
            "first_listing_line": start if projection["lines"] else None,
            "last_listing_line": len(listing) if projection["lines"] else None,
            **{key: value for key, value in projection.items() if key != "lines"},
        })
        total_effective += projection["effective_code_lines"]
        total_physical += projection["physical_lines"]
        total_statements += projection["logical_statement_lines"]
    page, occupied = [], 0
    for item in listing:
        needed = len(item["rows"])
        if needed > rows_per_page:
            raise ValueError(f"Source line exceeds one page: {item['file']}:{item['source_line']}")
        if page and occupied + needed > rows_per_page:
            pages.append(page)
            page, occupied = [], 0
        page.append(item)
        occupied += needed
    if page:
        pages.append(page)
    selected = list(range(len(pages))) if len(pages) <= 60 else list(range(30)) + list(range(len(pages) - 30, len(pages)))
    short_pages = [
        index + 1 for index, page in enumerate(pages)
        if sum(bool(item["text"].strip()) for item in page) < 50 and index != len(pages) - 1
    ]
    return {
        "files": files, "listing": listing, "pages": pages, "selected": selected,
        "totals": {
            "files": len(files), "physical_source_lines": total_physical,
            "effective_code_lines": total_effective,
            "logical_statement_lines": total_statements,
            "listing_lines": len(listing), "full_pages_planned": len(pages),
            "submitted_pages_planned": len(selected),
        },
        "nonfinal_pages_below_50_code_lines": short_pages,
    }


def _set_run_font(run, latin: str, size: float, chinese: str = "宋体") -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = latin
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:eastAsia"), chinese)


def _field(paragraph, name: str, cached: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    _set_run_font(run, "Arial", 8)
    run.bold = True
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" {name} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = cached
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, value, end):
        run._element.append(element)


def write_document(plan: dict, target: Path, software_name: str, version: str, *, full: bool = False) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(18)
    section.top_margin = section.bottom_margin = Mm(20)
    section.header_distance = section.footer_distance = Mm(10)
    normal = document.styles["Normal"]
    normal.font.name = "Arial Narrow"
    normal.font.size = Pt(FONT_SIZE)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(ROW_HEIGHT)
    normal.paragraph_format.widow_control = False
    normal.paragraph_format.keep_with_next = False
    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    _set_run_font(header.add_run(f"{software_name} "), "Arial", 10)
    header_version = header.add_run(version)
    _set_run_font(header_version, "Arial", 10)
    header_version.bold = True
    indices = list(range(len(plan["pages"]))) if full else plan["selected"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _field(footer, "PAGE", "1")
    _set_run_font(footer.add_run("of"), "Arial", 8)
    _field(footer, "NUMPAGES", str(len(indices)))
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    document.core_properties.title = f"{software_name} {version} 源程序"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    for page_index, full_index in enumerate(indices):
        for line_index, item in enumerate(plan["pages"][full_index]):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(LINE_NUMBER_WIDTH)
            paragraph.paragraph_format.first_line_indent = Pt(-LINE_NUMBER_WIDTH)
            paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(LINE_NUMBER_WIDTH))
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.page_break_before = bool(page_index and line_index == 0)
            _set_run_font(paragraph.add_run(str(item["number"]) + "\t"), "Times New Roman", FONT_SIZE)
            for index, fragment in enumerate(item["rows"]):
                if index:
                    paragraph.add_run().add_break()
                _set_run_font(paragraph.add_run(fragment), "Arial Narrow", FONT_SIZE)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)


def manifest(plan: dict, name: str, version: str, rows_per_page: int) -> dict:
    pages = []
    for index, page in enumerate(plan["pages"]):
        pages.append({
            "full_page": index + 1,
            "submitted_page": plan["selected"].index(index) + 1 if index in plan["selected"] else None,
            "first_listing_line": page[0]["number"], "last_listing_line": page[-1]["number"],
            "code_lines": sum(bool(item["text"].strip()) for item in page),
            "display_rows": sum(len(item["rows"]) for item in page),
        })
    return {
        "schema_version": 1, "software_name": name, "version": version,
        "source_order": [item["path"] for item in plan["files"]],
        "files": plan["files"], "totals": plan["totals"],
        "page_selection": "all" if len(plan["pages"]) <= 60 else "first_30_and_last_30",
        "pages": pages,
        "formatting_projection": {
            "description": "鉴别材料排版投影。原始程序未改写。仅移除可安全识别的注释、独立文档字符串与非字符串内部空行，制表符显示为4个空格，长行使用软换行。",
            "authorship": "材料排版与哈希不证明人工编写，也不改变实际开发过程。",
            "line_numbers": "按全量源程序顺序跨文件连续编号，前后30页选取时保留原始全量行号。",
            "source_text_mapping": "listing_map记录每个排版行对应的原文件行号，files哈希针对未经处理的原文件字节。",
            "code_font": "Arial Narrow 10.5pt / 宋体 10.5pt",
            "line_number_font": "Times New Roman 10.5pt",
            "line_spacing_pt": ROW_HEIGHT, "maximum_display_rows_per_page": rows_per_page,
            "render_verification": "planned_only; must confirm full and submitted page counts using Word/PDF rendering before delivery",
        },
        "excluded": {
            "tests": "测试代码不计入运行程序量",
            "third_party": "虚拟环境、依赖库、FFprobe程序、打包产物不计入自有源码",
            "data": "视频、截图、示例数据库、报告和用户数据均不计入",
            "tools": "除被运行入口直接引用的capture_materials.py和packaged_smoke.py外，构建、导出、验证工具不计入",
            "docs": "说明书、README与验证记录不计入程序量",
        },
        "nonfinal_pages_below_50_code_lines": plan["nonfinal_pages_below_50_code_lines"],
        "listing_map": [
            {"number": line["number"], "file": line["file"], "source_line": line["source_line"]}
            for line in plan["listing"]
        ],
    }


def clean_project_files(root: Path) -> list[Path]:
    """Select source, tests, build scripts and text documentation, never runtime data."""
    root = root.resolve()
    files = set(source_files(root))
    for directory in ("tests", "tools"):
        files.update((root / directory).rglob("*.py"))
    for suffix in ("*.md", "*.txt", "*.json"):
        files.update((root / "docs").rglob(suffix))
    for suffix in ("*.yml", "*.yaml"):
        files.update((root / ".github" / "workflows").rglob(suffix))
    files.update(root / name for name in PROJECT_FILES if (root / name).is_file())
    safe = []
    for path in files:
        if not path.is_file():
            continue
        relative = path.relative_to(root)
        if any(part.startswith(".") for part in relative.parts[:-1]) and relative.parts[0] != ".github":
            continue
        if path.is_symlink() or not path.resolve().is_relative_to(root):
            raise ValueError(f"Source package cannot include a link outside the project: {relative}")
        if any(part in {"__pycache__", "data", "reports", "screenshots", "build", "dist"} for part in relative.parts):
            continue
        safe.append(path)
    return sorted(safe, key=lambda path: path.relative_to(root).as_posix())


def write_source_zip(root: Path, plan: dict, target: Path) -> dict:
    root = root.resolve()
    contents = {path.relative_to(root).as_posix(): path.read_bytes() for path in clean_project_files(root)}
    for item in plan["files"]:
        if hashlib.sha256(contents[item["path"]]).hexdigest() != item["sha256"]:
            raise ValueError(f"Source changed after document planning: {item['path']}")
    package_manifest = {
        "description": "完整项目源码、测试、构建工具及文本配置，运行源文件保留原始注释。",
        "excluded": ["用户数据", "视频及图片", "报告", "数据库", "第三方依赖", "虚拟环境", "build", "dist", ".git"],
        "files": [
            {"path": name, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()}
            for name, data in contents.items()
        ],
    }
    contents["SOURCE_PACKAGE_MANIFEST.json"] = (
        json.dumps(package_manifest, ensure_ascii=False, indent=2) + "\n"
    ).encode("utf-8")
    target.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(target, "w", ZIP_DEFLATED) as archive:
        for relative, data in contents.items():
            entry = ZipInfo(relative, date_time=(2026, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            archive.writestr(entry, data)
    return package_manifest


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="生成可追溯的源程序鉴别材料，保留原始源码。")
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--name", default="智能会议设备投屏质量分析系统")
    parser.add_argument("--version", default="V1.0")
    parser.add_argument("--rows-per-page", type=int, default=ROWS_PER_PAGE)
    parser.add_argument("--plan-only", action="store_true")
    parser.add_argument("--include-full", action="store_true")
    parser.add_argument("--source-zip", action="store_true")
    args = parser.parse_args(argv)
    plan = build_plan(args.root.resolve(), font_measurement(), args.rows_per_page)
    result = manifest(plan, args.name, args.version, args.rows_per_page)
    if not args.plan_only:
        if args.output_dir is None:
            parser.error("--output-dir is required unless --plan-only is used")
        if plan["nonfinal_pages_below_50_code_lines"]:
            parser.error("长源代码行导致非末页少于50代码行，请调整源码本身的合理换行后重新生成。")
        args.output_dir.mkdir(parents=True, exist_ok=True)
        write_document(plan, args.output_dir / "源程序鉴别材料.docx", args.name, args.version)
        if args.include_full:
            write_document(plan, args.output_dir / "源程序全量排版.docx", args.name, args.version, full=True)
        if args.source_zip:
            package = write_source_zip(args.root, plan, args.output_dir / "完整项目源码.zip")
            result["source_package"] = {"filename": "完整项目源码.zip", "file_count": len(package["files"])}
        (args.output_dir / "source_manifest.json").write_text(
            json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
    print(json.dumps({
        **plan["totals"], "nonfinal_pages_below_50_code_lines": plan["nonfinal_pages_below_50_code_lines"],
        "files": [{key: item[key] for key in ("path", "effective_code_lines", "md5")} for item in plan["files"]],
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
