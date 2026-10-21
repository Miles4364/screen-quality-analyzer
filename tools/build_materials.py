"""使用已核验的源码清单和真实窗口截图生成 V1.0 使用说明及信息表。

需要 python-docx、Pillow。先生成 source_manifest.json，再运行此工具。
此工具不代替文档渲染检查，也不自动填写权利人或开发完成日期。
"""
import argparse
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from utils.product import APP_NAME, DISPLAY_VERSION


FUNCTIONS = (
    '本系统面向会议设备投屏录制视频的质量检查，采用 Python、PySide6 和 SQLite 构建本地桌面应用，'
    '把视频读取、异常定位、画面复核和结果归档组织为可重复执行的操作流程。分析工作台采用可调整宽度的三栏布局，'
    '左侧按项目展示近期任务，中间显示分析进度与异常列表，右侧显示视频属性及选中事件详情。用户选择本地视频后，'
    '可指定项目和任务名称；后台读取分辨率、帧率、编码及容器时长，随后逐帧解码，避免一次性加载整段视频。'
    '检测过程将画面缩放至 320×180 并转换为灰度图，以平均亮度低于 10 的连续片段识别黑屏候选，以相邻帧中'
    '灰度差超过 5 的像素比例低于 0.01 识别疑似冻结候选，两类候选均需持续至少 0.5 秒才形成事件。'
    '黑屏候选会中断冻结候选，事件保留起止时间、持续时间和严重等级。对于容器帧数与解码帧数的极小差异，'
    '系统调用 FFprobe 独立计数核验；明显截断、核验超时和用户取消仍按未完成处理。结果保存到本地数据库后，'
    '用户可以打开异常附近画面、逐帧移动、播放片段并保存截图，也可以将事件标记为待复核、确认异常或误报，'
    '记录备注并按状态筛选。统计图表页面展示异常时间轴、事件持续时间散点、已分析时长构成和人工复核数量，'
    '所有图形共享当前任务数据。时长构成对重叠区间去重，重叠部分优先计入黑屏，辅助评分按黑屏和疑似冻结的'
    '时长占比分别扣分，人工复核状态不改写自动检测和评分。历史记录支持查找任务、修改名称和调整项目归属，'
    '原视频缺失时仍能阅读已保存的事件及截图。Excel 报告汇总元数据、事件、评分规则与复核记录，命令行入口'
    '提供同一分析引擎的文本或 JSON 输出。当前时间定位以固定帧率为前提；暗场、静态演示文稿及低运动画面可能'
    '触发候选，输出用于辅助排查，不能仅凭事件数量或分数判断设备故障，也不能将合成样本结果解释为真实会议准确率。'
)


def build(output, screenshots):
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from PIL import Image, ImageDraw, ImageFont

    manifest = json.loads((output / 'source_manifest.json').read_text(encoding='utf-8'))
    capture = json.loads((screenshots / 'capture.json').read_text(encoding='utf-8'))
    if not capture.get('ok'):
        raise ValueError('截图流程尚未成功，请先完成可见桌面截图')
    if capture['statistics'] != {'black_count': 2, 'freeze_count': 2, 'total_count': 4}:
        raise ValueError('材料示例需要 60 秒标准合成验证片，请检查实际分析输出')
    if capture['score'] != 90.83 or capture['processed_frames'] != 1800:
        raise ValueError('示例结果与标准验证片不匹配')
    refs = json.loads((ROOT / 'docs/materials-v1/references.json').read_text(encoding='utf-8'))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(18)
    section.top_margin = section.bottom_margin = Mm(20)
    section.header_distance = section.footer_distance = Mm(10)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption'):
        style = doc.styles[name]
        style.font.name = '宋体'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
        for border in style.element.findall('.//' + qn('w:pBdr')):
            border.getparent().remove(border)
    normal = doc.styles['Normal']
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(23)
    normal.paragraph_format.first_line_indent = Pt(24)
    for name, size in (('Heading 1', 22), ('Heading 2', 12)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(8)
    doc.core_properties.title = APP_NAME + ' 使用说明 ' + DISPLAY_VERSION
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''

    def font(run, size=12, latin='宋体', bold=False, italic=False):
        run.font.name = latin
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = RGBColor(0, 0, 0)
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
        return run

    def field(p, name, bold=False):
        run = font(p.add_run(), 8, 'Arial', bold)
        f = OxmlElement('w:fldSimple')
        f.set(qn('w:instr'), name)
        r = OxmlElement('w:r')
        r.append(run._element.get_or_add_rPr())
        t = OxmlElement('w:t')
        t.text = '1'
        r.append(t)
        f.append(r)
        p._p.append(f)

    header = section.header.paragraphs[0]
    header.paragraph_format.first_line_indent = Pt(0)
    font(header.add_run(APP_NAME + '使用说明'), 10)
    font(header.add_run(DISPLAY_VERSION), 10, 'Arial', True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.first_line_indent = Pt(0)
    field(footer, 'PAGE', True)
    font(footer.add_run('of'), 8, 'Arial')
    field(footer, 'NUMPAGES', True)

    def para(text, size=12, spacing=23, indent=True):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = Pt(spacing)
        p.paragraph_format.first_line_indent = Pt(24 if indent else 0)
        font(p.add_run(text), size)
        return p

    def heading(text, level=2):
        p = doc.add_paragraph(style=f'Heading {level}')
        number, _, title = text.partition(' ')
        if title and number[0].isdigit():
            font(p.add_run(number + ' '), 22 if level == 1 else 12,
                 'Calibri' if level == 1 else 'Arial', True)
            font(p.add_run(title), 22 if level == 1 else 12, bold=True)
        else:
            font(p.add_run(text), 22 if level == 1 else 12, bold=True)
        return p

    def citation(p, numbers):
        run = font(p.add_run(''.join(f'[{n}]' for n in numbers)), 9)
        run.font.superscript = True

    def new_page(title, level=2):
        doc.add_page_break()
        heading(title, level)

    def figure(filename, chinese, english, index, width=174):
        path = screenshots / (filename + '.png')
        if not path.is_file():
            raise ValueError(f'缺少真实截图：{path}')
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        with Image.open(path) as picture:
            width = min(width, 88 * picture.width / picture.height)
        p.add_run().add_picture(str(path), width=Mm(width))
        caption = doc.add_paragraph()
        caption.paragraph_format.first_line_indent = Pt(0)
        caption.paragraph_format.line_spacing = Pt(12)
        caption.paragraph_format.space_after = Pt(3)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(caption.add_run(f'图 3-{index} {chinese}（{filename}.png）\n'), 9)
        font(caption.add_run(f'Figure 3-{index} {english}'), 9, 'Times New Roman')

    # 封面使用独立段落排列，避免把竖排标题做成不可编辑的图片。
    para('', spacing=36)
    title = doc.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.space_after = Pt(22)
    borders = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right', 'between', 'bar'):
        border = OxmlElement('w:' + edge)
        border.set(qn('w:val'), 'nil')
        borders.append(border)
    title._p.get_or_add_pPr().append(borders)
    font(title.add_run(APP_NAME), 28, bold=True)
    p = para('版本 1.0', 16, 28, False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para('', spacing=36)
    p = para('用\n户\n手\n册', 28, 42, False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para('', spacing=40)
    p = para('二〇二六年10月', 16, 25, False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    new_page('1 功能描述', 1)
    p = para(FUNCTIONS, 12, 20)
    citation(p, [1, 2, 3, 6, 7, 8])
    refmark = font(p.add_run(), 9)
    refmark.font.superscript = True
    note = OxmlElement('w:footnoteReference')
    note.set(qn('w:id'), '1')
    refmark._element.append(note)
    footnote_xml = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
 <w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
 <w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
 <w:footnote w:id="1"><w:p><w:pPr><w:spacing w:line="200"/></w:pPr>
 <w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>
 <w:r><w:rPr><w:rFonts w:ascii="宋体" w:eastAsia="宋体"/><w:sz w:val="18"/></w:rPr>
 <w:t>参考文献说明所用接口；检测阈值与 duration-v1 评分是本项目的实现参数，不是上述文档提供的行业标准。</w:t></w:r>
 </w:p></w:footnote></w:footnotes>'''
    part = Part(PackURI('/word/footnotes.xml'),
                'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml',
                footnote_xml.encode('utf-8'), doc.part.package)
    doc.part.relate_to(part, RT.FOOTNOTES)

    new_page('1.1 判定方法与适用边界')
    para('下列计算对应当前源程序。灰度运算分别使用 OpenCV 和 NumPy 的图像差分、均值接口；阈值来自项目配置。', spacing=21)
    citation(doc.paragraphs[-1], [4, 5])
    formulas = [
        ('Bₖ = (1 / M) ∑ Gₖ(p)', '其中，M＝320×180，Gₖ(p) 为第 k 帧第 p 个灰度像素。Bₖ＜10 为黑屏候选，持续至少 0.5 秒形成事件。'),
        ('Rₖ = (1 / M) ∑ 1(|Gₖ(p) − Gₖ₋₁(p)| > 5)', '其中，1(条件) 在条件成立时取 1，否则取 0。Rₖ＜0.01 为相似帧对，持续至少 0.5 秒形成疑似冻结事件，起点包含前一帧。'),
        ('S = round(clamp(100 − 100B / T − 50F / T, 0, 100), 2)', '其中，T 为实际读取帧数除以有效帧率；B 为黑屏区间并集时长；F 为冻结区间并集扣除黑屏重叠后的时长。clamp 限定分数在 0 至 100，round 保留两位小数。'),
    ]
    for expression, explanation in formulas:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(6)
        font(p.add_run(expression), 12, 'Times New Roman', italic=True)
        para(explanation, spacing=21)
    para('严重等级按持续时间划分：不足 1 秒为低，1 秒至不足 2 秒为中，达到 2 秒为高。固定帧率下采用左闭右开的事件时间区间。变帧率视频未逐帧使用显示时间戳，不能承诺其定位精度。')
    para('无有效分析时长、帧数或事件时间无效时不输出有效评分。暗场可能被判为黑屏，静态 PPT 可能被判为疑似冻结，应查看异常前后画面再填写人工复核结论。')

    new_page('2 项目结构', 1)
    para('下列运行程序按源程序材料的实际拼接顺序列出。空的包初始化文件保留在项目中，不产生代码材料行号；测试、构建工具和第三方依赖不计入运行程序量。完整逐文件哈希与行号映射见随附 source_manifest.json。', spacing=21)
    tree = ['screen-quality-analyzer/']
    previous = None
    for item in manifest['files']:
        path = Path(item['path'])
        parent = path.parent.as_posix()
        if parent != '.':
            if parent != previous:
                tree.append('├── ' + parent + '/')
            tree.append('│   ├── ' + path.name)
        else:
            tree.append('├── ' + path.name)
        previous = parent
    p = para('\n'.join(tree), 9.5, 12, False)
    for r in p.runs:
        r.font.name = 'Consolas'
    new_page('2.1 源码数量与版本核对')
    totals = manifest['totals']
    para(f'本次运行程序共 {totals["files"]} 个 Python 文件，物理行数为 {totals["physical_source_lines"]}，'
         f'去除可识别注释、独立文档字符串及非字符串空行后为 {totals["effective_code_lines"]} 个有效代码行。'
         f'全量排版计划为 {totals["full_pages_planned"]} 页，鉴别材料选取 {totals["submitted_pages_planned"]} 页。'
         '材料中的软换行不新增源码行号；哈希针对原始文件字节计算。')
    for filename in ('gui_main.py', 'core/analysis_engine.py', 'core/black_detector.py', 'core/freeze_detector.py', 'gui/main_window.py'):
        item = next(row for row in manifest['files'] if row['path'] == filename)
        para(filename + '\nMD5 ' + item['md5'], 10, 17, False)
    para('源程序材料仅做阅读排版投影，原项目保留注释、异常处理、数据类及开发记录。若全量排版超过 60 页，材料保留连续前 30 页与连续后 30 页，编号沿用全量顺序，中间编号跳跃对应明确省略区间。')
    para('程序与材料统一使用 V1.0。界面保持 PySide6；图形布局和文字说明按规范调整。没有为制造代码量而加入空语句、无关业务或第三方实现。')

    new_page('3 系统功能', 1)
    heading('3.1 启动与文件准备')
    para('Windows 便携版解压后运行 ScreenQualityAnalyzer.exe，保留整个目录及 tools 中的 FFprobe。程序默认最大化，可拖动三栏分隔线调整显示宽度。数据存入当前用户的应用数据目录，重新安装前应备份数据库及已保存截图。')
    para('源码方式先安装 requirements.txt 和 requirements-gui.txt，配置可用 FFprobe 后运行 python gui_main.py。FFprobe 查找顺序与覆盖选项见项目 README；SQA_FFPROBE 可以显式指定程序路径。')
    para('本轮核验环境为 Windows 11 x64、Python 3.11.2、Intel Core i7-11700F、16 GB 内存。此配置是实测环境，不代表最低配置。便携包携带解释器及依赖；Windows 10 和其他硬件尚需独立验收。', spacing=21)
    para('本手册截图使用确定性合成验证视频：1920×1080、30 FPS、60 秒、1800 帧。其注入区间为黑屏 10–12 秒及 40–41 秒，冻结 25–27 秒及 50–53 秒。本次实际解码识别出黑屏 2 段、疑似冻结 2 段，辅助评分 90.83。截图中的项目名、复核备注是流程验证记录。', spacing=21)
    para('常见问题处理：若提示找不到 FFprobe，确认解压目录完整或配置程序路径；若显示读取提前结束，应保留具体错误并核验源文件，不能仅靠改名或跳过检查得到成功结果；如视频已移动或修改，历史结果可以查看，但画面复核会受到限制。')
    para('分析大文件时主要开销来自解码和帧处理，应保证磁盘空间与视频可读性。点击“取消分析”后等待当前操作结束；取消的任务不会作为完成结果保存。统计得分是辅助指标，应结合原画面复核。')

    pages = [
        ('3.2 创建项目与选择视频',
         '点击“选择视频”导入本地录制文件，再选择所属项目。点击“新建项目”填写名称并确认；任务名称可自行填写，留空使用视频文件名。左侧显示最近 200 个任务，更多任务从“历史记录”查询。',
         [('xiangmu_01','输入新项目名称','Entering a project name'),('fenxi_01','工作台分析结果','Analysis workspace')]),
        ('3.3 分析结果与人工复核',
         '完成分析后结果自动保存，异常列表显示类型、起止时间、时长和等级。先选中一条异常，再点击“记录人工复核”。选择状态并填写备注，点击“保存复核结论”；此操作不修改自动检测区间。',
         [('rengong_01','打开复核记录','Opening the review record'),('rengong_02','设置复核状态和备注','Setting review status and notes')]),
        ('3.4 查看复核结果与筛选',
         '保存后，列表与右侧详情同步显示复核状态。使用复核状态下拉框仅显示“确认异常”等指定状态；切回“全部复核状态”恢复所有事件。筛选只改变显示范围，统计图与评分仍基于完整任务。',
         [('fenxi_02','保存后的复核状态','Saved review state'),('shaixuan_01','按确认异常筛选','Filtering confirmed events')]),
        ('3.5 定位异常与检查前后画面',
         '选中异常并点击“查看截图 / 跳转视频”，打开画面复核窗口。窗口显示实际解码帧及定位时间。点击“查看前 2 秒”跳转事件开始前的上下文画面，结合动态内容判断静止是否符合原素材。',
         [('huamian_01','异常起点画面','Frame at event start'),('huamian_02','异常前上下文','Context before the event')]),
        ('3.6 逐帧移动与截图记录',
         '使用上一帧、下一帧或片段播放检查变化。点击“保存截图”选择位置，截图会与任务和事件关联。返回工作台后点击“已保存截图”查看记录；保存的截图是视频帧，说明书中的图片则是运行界面的窗口截图。',
         [('huamian_03','移动到相邻帧','Moving to the next frame'),('jietu_01','已保存截图列表','Saved frame captures')]),
        ('3.7 异常时间轴与持续时间',
         '点击顶部“统计图表”，在标签页之间切换。时间轴横向表示视频秒数，按黑屏、疑似冻结分行；散点图横轴为事件起点，纵轴为持续秒数。两图展示当前任务实际检出的事件，空结果会明确提示未检出。',
         [('tongji_01','异常时间轴','Event timeline'),('tongji_02','事件持续时间散点图','Event duration scatter plot')]),
        ('3.8 时长构成与人工复核统计',
         '“时长构成”将实际分析时长分成黑屏、非黑屏疑似冻结和其余时长，区间重叠只计算一次。“人工复核”统计待复核、确认异常和误报数量。它们分别回答时间占用与复核进度，不能互相替代。',
         [('tongji_03','已分析时长构成','Analyzed duration composition'),('tongji_04','人工复核数量','Review status counts')]),
        ('3.9 历史查询与项目归档',
         '点击“历史记录”按任务名、视频文件名或项目查找保存的任务，列表显示保存时间。选中任务可改名或移动到其他项目，再打开历史结果继续查看。此例把同一条验证任务移入“复核样本归档”。',
         [('lishi_01','历史记录初始状态','History before organization'),('lishi_02','任务改名并调整项目','Renamed and regrouped task')]),
        ('3.10 打开历史结果与导出报告',
         '打开任务后可继续查看元数据、异常列表、图表及复核记录。点击“导出 Excel 报告”选择目标文件，完成后界面提示文件名。已有目标文件需确认覆盖；导出失败保留当前分析结果，可以更换位置重试。',
         [('lishi_03','重新打开归档任务','Reopening an archived task'),('daochu_01','报告导出完成','Completed report export')]),
    ]
    number = 1
    for title, body, pictures in pages:
        new_page(title)
        para(body, 11, 19)
        for filename, cn, en in pictures:
            figure(filename, cn, en, number)
            number += 1
    new_page('3.11 命令行与验证方法')
    para('命令行与桌面共用 AnalysisEngine。下列路径是示例，请替换为已有视频。命令不会上传视频；JSON 输出可用于后续数据处理。')
    for command in ('python main.py "samples/最终验收 1080p.mp4"',
                    'python main.py "samples/最终验收 1080p.mp4" --json',
                    'python -m pytest -q',
                    'python gui_main.py "samples/最终验收 1080p.mp4" --auto-screenshot'):
        p = para(command, 10, 18, False)
        p.runs[0].font.name = 'Consolas'
    para('自动截图需要可见的 Windows 桌面。它使用独立验证数据库，完成真实分析、复核状态保存、历史打开和报告导出，再调用系统窗口截屏；只产生单色截图时会失败。截图记录保存在 screenshots/v1/capture.json，不能将离屏绘制替换为真实桌面截图。')
    para('源码测试覆盖视频读取、黑屏和冻结阈值、异常边界、取消、数据库、报告、评分及工作台数据联动。通过测试并不等同于真实会议准确率评估；本轮最终结果和打包验证记录见项目 docs/V1.0-validation.md。')
    para('Windows 打包采用 PyInstaller 单目录方式，同时提供 GUI、CLI 和内置 FFprobe。测试记录区分源码测试、本机隔离 PATH 验证与独立机器验收，后者不能由本机打包成功替代。', spacing=21)
    citation(doc.paragraphs[-1], [9])

    new_page('参考文献', 1)
    for ref in refs:
        para(f'[{ref["id"]}] {ref["author"]}. {ref["title"]}[EB/OL]. '
             f'[{ref["accessed"]}]. {ref["url"]}.', 10, 16, False)
    para('上述文献支持接口和技术机制说明；实现参数、评分权重、界面组织与验证结果以本项目源程序及实际记录为准。', 10, 17)

    # 原生黑白流程图用于结构说明，不替代任何软件截图。
    image = Image.new('RGB', (1500, 1500), 'white')
    draw = ImageDraw.Draw(image)
    font_path = 'C:/Windows/Fonts/simsun.ttc'
    chart_font = ImageFont.truetype(font_path, 32)
    boxes = [(450,40,1050,160,'本地视频与分析配置'),
             (450,230,1050,350,'FFprobe 元数据核验'),
             (450,420,1050,540,'VideoReader 逐帧读取'),
             (80,620,700,760,'黑屏检测'),(800,620,1420,760,'疑似冻结检测'),
             (450,850,1050,990,'AnalysisEngine 事件与评分'),
             (80,1090,700,1230,'SQLite 保存任务与复核'),(800,1090,1420,1230,'工作台与四类统计图'),
             (450,1330,1050,1460,'画面复核与 Excel 报告')]
    for left,top,right,bottom,text in boxes:
        draw.rectangle((left,top,right,bottom),outline='black',width=3)
        draw.text(((left+right)/2,(top+bottom)/2),text,font=chart_font,anchor='mm',fill='black')
    def arrow(a,b):
        draw.line((a,b),fill='black',width=3)
        x,y=b
        draw.polygon(((x,y),(x-9,y-18),(x+9,y-18)),fill='black')
    for a,b in [((750,160),(750,230)),((750,350),(750,420)),
                ((750,540),(390,620)),((750,540),(1110,620)),
                ((390,760),(650,850)),((1110,760),(850,850)),
                ((750,990),(390,1090)),((750,990),(1110,1090)),
                ((390,1230),(650,1330)),((1110,1230),(850,1330))]:
        arrow(a,b)
    diagram = output / '系统处理流程图.png'
    image.save(diagram)
    new_page('系统处理流程', 1)
    para('界面与命令行调用同一分析引擎。数据库保存结构化结果和人工复核，图表由已保存或当前完成的结果计算；视频帧按需读取。流程图只描述当前实现，不包含联网识别或模型训练。')
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(diagram), width=Mm(170))
    doc.save(output / '软件使用说明书.docx')

    fields = {
        '软件名称': APP_NAME, '版本号': DISPLAY_VERSION,
        '开发的硬件环境': 'Intel Core i7-11700F @ 2.50 GHz，16 GB 内存；本轮实测主机',
        '运行的硬件环境': '已验证上述 x64 主机；最低硬件配置尚未测定',
        '开发该软件的操作系统': 'Windows 11 家庭中文版 x64，10.0.26200',
        '软件开发环境 / 开发工具': 'Python 3.11.2，Git，pytest，PyInstaller；当前项目使用 AI 辅助开发',
        '该软件的运行平台 / 操作系统': '已验证 Windows 11 x64；其他环境待独立验证',
        '软件运行支撑环境 / 支持软件': 'Python 3.11.2，PySide6 6.12.0，OpenCV 4.14.0.94，NumPy 2.4.6，openpyxl 3.1.5，FFprobe 5.1.2，SQLite；便携包包含运行依赖',
        '编程语言': 'Python',
        '源程序量': f'{totals["effective_code_lines"]} 行有效代码（{totals["physical_source_lines"]} 物理行，{totals["files"]} 个运行文件；不含测试和第三方代码）',
        '鉴别材料范围': f'全量排版 {totals["full_pages_planned"]} 页，材料 {totals["submitted_pages_planned"]} 页；详细行号和哈希见 source_manifest.json',
        '开发目的': '辅助会议设备测试人员定位录制视频中的黑屏和疑似冻结，记录人工复核并导出可追溯报告。',
        '面向领域': '会议设备测试', '软件的主要功能': FUNCTIONS,
        '软件的技术特点': '本地逐帧检测与结果追溯',
        '填报说明': '以上为项目技术信息。权利人、权利归属、开发完成日期和发表情况由申请人核实，材料格式不证明代码由人工独立编写。',
    }
    (output / '软件信息.txt').write_text(''.join(f'{k}：{v}\n' for k,v in fields.items()), encoding='utf-8')
    (output / 'manual_manifest.json').write_text(json.dumps({
        'name': APP_NAME, 'version': DISPLAY_VERSION, 'function_characters': len(FUNCTIONS),
        'screenshots': [name for _,_,images in pages for name,_,_ in images],
        'sample': capture, 'references': refs,
    }, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'manual': '软件使用说明书.docx', 'function_characters': len(FUNCTIONS),
                      'figures': number-1}, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--screenshots', type=Path, default=ROOT/'screenshots/v1')
    args = parser.parse_args()
    build(args.output_dir.resolve(), args.screenshots.resolve())
