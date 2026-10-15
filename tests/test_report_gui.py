import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import time
import pytest
pytest.importorskip('PySide6')
from openpyxl import load_workbook
from tests.test_gui import window, wait_until
from tests.test_review import result
from reporting.excel_report import ReportCancelled


def test_gui_export_saved_history(app,window,result,tmp_path):
    window._show_result(result)
    task_id=window.task_id
    window.store.rename_task(task_id,'历史任务名')
    window.load_history(task_id)
    target=tmp_path / 'GUI.xlsx'
    window.export_report_to(target)
    assert not window.start_button.isEnabled()
    wait_until(app,lambda:window.worker is None)
    assert target.is_file()
    assert window.result is not None
    assert window.export_button.isEnabled()
    book=load_workbook(target)
    summary=dict(book['分析概况'].iter_rows(min_row=2,values_only=True))
    assert summary['任务名称']=='历史任务名'
    book.close()
    window.export_report_to(target)
    wait_until(app,lambda:window.worker is None)
    assert '导出失败' in window.status_label.text()
    assert window.result is not None


def test_cancel_report_keeps_analysis(app,window,result,tmp_path,monkeypatch):
    window._show_result(result)
    def slow(*args,should_cancel,**kwargs):
        end=time.monotonic()+5
        while time.monotonic()<end:
            if should_cancel():
                raise ReportCancelled()
            time.sleep(.005)
        raise RuntimeError('cancel not delivered')
    monkeypatch.setattr('gui.report_worker.export_report',slow)
    target=tmp_path / 'cancel.xlsx'
    window.export_report_to(target)
    window.cancel_analysis()
    wait_until(app,lambda:window.worker is None)
    assert window.result is result
    assert not target.exists()
    assert '已取消' in window.status_label.text()
    assert window.export_button.isEnabled()


def test_export_unsaved_result_when_database_unavailable(app,window,result,tmp_path,monkeypatch):
    from storage.analysis_store import StorageError
    def fail(*args,**kwargs):
        raise StorageError('database locked')
    monkeypatch.setattr(window.store,'save_result',fail)
    window._show_result(result)
    assert window.task_id is None
    target=tmp_path / 'unsaved.xlsx'
    window.export_report_to(target)
    wait_until(app,lambda:window.worker is None)
    assert target.exists()
    assert window.save_result_button.isEnabled()
