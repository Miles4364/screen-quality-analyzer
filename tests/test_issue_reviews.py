from dataclasses import asdict
import sqlite3
import pytest
from openpyxl import load_workbook
from tests.test_storage import result
from core.review import review_counts
from storage.analysis_store import AnalysisStore, StorageError
from reporting.excel_report import export_report


def test_review_roundtrip_report_and_score_unchanged(tmp_path, result):
    store = AnalysisStore(tmp_path/'db.sqlite3')
    task_id = store.save_result(result)
    before, score = asdict(result), result.quality
    assert store.load_task(task_id)['reviews'] == []
    assert review_counts(1, [])['PENDING'] == 1
    store.save_review(task_id, 0, 'FALSE_POSITIVE', '=中文备注\n静态 PPT')
    task = AnalysisStore(store.path).load_task(task_id)
    assert task['reviews'][0]['note'] == '=中文备注\n静态 PPT'
    assert task['reviews'][0]['updated_at']
    assert asdict(task['result']) == before
    assert task['result'].quality == score
    book = load_workbook(export_report(result, tmp_path/'report.xlsx', task=task))
    assert book['异常明细']['K2'].value == '误报'
    assert book['异常明细']['L2'].value == '=中文备注\n静态 PPT'
    assert book['异常明细']['L2'].data_type == 's'
    assert book['异常明细']['M2'].value == task['reviews'][0]['updated_at']
    assert dict(book['分析概况'].iter_rows(min_row=2, values_only=True))['人工复核：误报'] == 1
    book.close()
    store.save_review(task_id, 0, 'CONFIRMED', '确认黑屏')
    store.save_review(task_id, 0, 'PENDING', '')
    assert len(store.load_task(task_id)['reviews']) == 1
    assert review_counts(1, store.load_task(task_id)['reviews'])['PENDING'] == 1


@pytest.mark.parametrize('ordinal,status,note', [(9,'PENDING',''),(-1,'PENDING',''),(0,'BAD',''),(0,'PENDING','x'*2001),(0,'PENDING',None)])
def test_invalid_review_does_not_change_data(tmp_path, result, ordinal, status, note):
    store = AnalysisStore(tmp_path/'db.sqlite3')
    task_id = store.save_result(result)
    store.save_review(task_id,0,'CONFIRMED','原备注')
    before = store.load_task(task_id)['reviews']
    with pytest.raises(StorageError):
        store.save_review(task_id,ordinal,status,note)
    assert store.load_task(task_id)['reviews'] == before


def test_v2_upgrade_and_rollback(tmp_path, result, monkeypatch):
    store = AnalysisStore(tmp_path/'db.sqlite3')
    task_id = store.save_result(result)
    with sqlite3.connect(store.path) as db:
        db.execute('DROP TABLE reviews')
        db.execute('PRAGMA user_version=2')
    original = AnalysisStore._migrate_v3
    def fail(db):
        original(db)
        raise sqlite3.OperationalError('migration failure')
    monkeypatch.setattr(AnalysisStore, '_migrate_v3', staticmethod(fail))
    with pytest.raises(StorageError):
        store.load_task(task_id)
    with sqlite3.connect(store.path) as db:
        assert db.execute('PRAGMA user_version').fetchone()[0] == 2
        assert not db.execute("SELECT name FROM sqlite_master WHERE name='reviews'").fetchall()
    monkeypatch.setattr(AnalysisStore, '_migrate_v3', staticmethod(original))
    task = store.load_task(task_id)
    assert asdict(task['result']) == asdict(result)
    assert task['reviews'] == []
    with sqlite3.connect(store.path) as db:
        assert db.execute('PRAGMA user_version').fetchone()[0] == 3


def test_locked_review_preserves_previous(tmp_path,result):
    store = AnalysisStore(tmp_path/'db.sqlite3')
    task_id = store.save_result(result)
    store.save_review(task_id,0,'CONFIRMED','原备注')
    with sqlite3.connect(store.path) as db:
        db.execute('BEGIN IMMEDIATE')
        with pytest.raises(StorageError):
            store.save_review(task_id,0,'FALSE_POSITIVE','新备注')
    assert store.load_task(task_id)['reviews'][0]['note'] == '原备注'
