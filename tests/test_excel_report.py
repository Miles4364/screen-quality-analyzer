from dataclasses import asdict
import hashlib
import json
from pathlib import Path
import pytest
from openpyxl import load_workbook
from reporting.excel_report import export_report, ReportError, ReportCancelled
from storage.analysis_store import AnalysisStore
from tests.test_storage import result


def overview(book):
    return dict(book['分析概况'].iter_rows(min_row=2, values_only=True))


def test_report_saved_task_roundtrip(tmp_path, result):
    store = AnalysisStore(tmp_path / 'db.sqlite3')
    project = store.create_project('会议室 A')
    task_id = store.save_result(result, title='中文报告', project_id=project)
    screenshot = tmp_path / '截图.png'
    screenshot.write_bytes(b'fixture')
    store.add_screenshot(task_id,0,screenshot,30,1)
    task = store.load_task(task_id)
    original = asdict(result)
    target = export_report(result,tmp_path / '中文 报告.xlsx',task=task)
    book = load_workbook(target)
    assert book.sheetnames == ['分析概况','异常明细','截图关联']
    summary = overview(book)
    assert summary['任务名称'] == '中文报告'
    assert summary['所属项目'] == '会议室 A'
    assert summary['异常总数'] == 1
    assert summary['关联截图数'] == 1
    assert book['异常明细']['E2'].value == 1
    assert book['异常明细']['G2'].value == 1
    assert book['截图关联']['E2'].value == '可用'
    assert book['截图关联']['F2'].hyperlink.target == screenshot.as_uri()
    assert book['异常明细'].freeze_panes == 'A2'
    assert book['异常明细'].auto_filter.ref == 'A1:J2'
    assert asdict(result) == original
    book.close()


def test_empty_report_and_input_text_is_not_formula(tmp_path, result):
    result.issues = []
    result.video_metadata.filename = '=HYPERLINK("https://example.invalid")'
    path = export_report(result,tmp_path / 'zero.xlsx',task={'title':'=1+1\x00'})
    book = load_workbook(path)
    assert overview(book)['异常总数'] == 0
    assert overview(book)['任务名称'] == '=1+1'
    assert book['异常明细']['B2'].value == '未检测到异常'
    assert book['截图关联']['E2'].value == '无已关联截图'
    assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)
    book.close()


@pytest.mark.parametrize('state',['missing','changed'])
def test_unavailable_screenshot_has_no_link(tmp_path,result,state):
    path = tmp_path / 's.png'
    if state == 'changed':
        path.write_bytes(b'changed')
    shot = {'path':str(path),'sha256':hashlib.sha256(b'original').hexdigest(),
            'issue_ordinal':0,'frame_index':30,'timestamp':1,'created_at':'2026-10-15'}
    report = export_report(result,tmp_path / 'report.xlsx',screenshots=[shot])
    book = load_workbook(report)
    assert book['截图关联']['E2'].value != '可用'
    assert book['截图关联']['F2'].hyperlink is None
    book.close()


def test_overwrite_and_atomic_failure(tmp_path,result,monkeypatch):
    path = tmp_path / 'report.xlsx'
    path.write_bytes(b'previous report')
    with pytest.raises(ReportError,match='已存在'):
        export_report(result,path)
    import reporting.excel_report as module
    def fail(self,path):
        Path(path).write_bytes(b'partial')
        raise OSError('disk full')
    monkeypatch.setattr(module.Workbook,'save',fail)
    with pytest.raises(ReportError,match='disk full'):
        export_report(result,path,overwrite=True)
    assert path.read_bytes() == b'previous report'
    assert not list(tmp_path.glob('.report-*.xlsx'))


def test_cancel_before_publish_keeps_existing(tmp_path,result,monkeypatch):
    import reporting.excel_report as module
    path = tmp_path / 'report.xlsx'
    path.write_bytes(b'keep')
    cancelled = []
    original = module.Workbook.save
    def save(self,target):
        original(self,target)
        cancelled.append(True)
    monkeypatch.setattr(module.Workbook,'save',save)
    with pytest.raises(ReportCancelled):
        export_report(result,path,overwrite=True,should_cancel=lambda: bool(cancelled))
    assert path.read_bytes() == b'keep'
    assert not list(tmp_path.glob('.report-*.xlsx'))


def test_protected_path_and_extension(tmp_path,result):
    path = tmp_path / 'protected.xlsx'
    path.write_bytes(b'database')
    with pytest.raises(ReportError,match='不能覆盖'):
        export_report(result,path,overwrite=True,protected_paths=[path])
    with pytest.raises(ReportError,match='扩展名'):
        export_report(result,tmp_path / 'report.csv')
    assert path.read_bytes() == b'database'


def test_explicit_overwrite_and_cli(tmp_path,result):
    report = tmp_path / 'report.xlsx'
    export_report(result,report)
    result.issues=[]
    export_report(result,report,overwrite=True)
    with report.open('rb') as stream:
        book=load_workbook(stream)
        assert overview(book)['异常总数']==0
        book.close()
    import shutil, subprocess, sys
    from tools.generate_samples import generate
    if not shutil.which('ffprobe'):
        pytest.skip('需要 FFprobe')
    video=generate(tmp_path / 'cli.mp4',duration=1,black=(),freeze=())
    target=tmp_path / 'cli.xlsx'
    run=subprocess.run([sys.executable,'main.py',str(video),'--json','--report',str(target)],capture_output=True)
    assert run.returncode==0,run.stderr
    assert Path(json.loads(run.stdout)['report_path'])==target
    fail=subprocess.run([sys.executable,'main.py',str(video),'--report',str(target)],capture_output=True)
    assert fail.returncode==1 and b'Traceback' not in fail.stderr
