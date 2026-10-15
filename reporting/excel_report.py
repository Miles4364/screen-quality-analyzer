"""生成静态 XLSX 报告；不执行输入文本，不修改源数据。"""
from dataclasses import asdict
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import tempfile
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from core.models import AnalysisConfig
from utils.time_util import seconds_to_time


class ReportError(Exception):
    pass


class ReportCancelled(ReportError):
    pass


def _put(ws, row, column, value):
    cell = ws.cell(row, column)
    if isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub('', value)
        # 即使任务名以 = 开头，也只保存为字符串，不生成公式。
        cell.data_type = 's'
    else:
        cell.value = value
    cell.font = Font(name='Microsoft YaHei', size=11, color='24344A')
    cell.alignment = Alignment(vertical='top', wrap_text=True)
    if isinstance(value, float):
        cell.number_format = '0.000'
    return cell


def _table(workbook, title, headers, widths):
    ws = workbook.create_sheet(title)
    ws.sheet_view.showGridLines = False
    for column, (header, width) in enumerate(zip(headers, widths), 1):
        cell = _put(ws, 1, column, header)
        cell.fill = PatternFill('solid', fgColor='245FCC')
        cell.font = Font(name='Microsoft YaHei', size=11, color='FFFFFF', bold=True)
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:1'
    return ws


def export_report(result, output, *, config=None, task=None, screenshots=None,
                  overwrite=False, protected_paths=(), should_cancel=None):
    """返回实际导出路径；目标冲突/失败不破坏原文件。"""
    def check_cancel():
        if should_cancel is not None and should_cancel():
            raise ReportCancelled('报告导出已取消')

    output = Path(output).resolve()
    if output.suffix.lower() != '.xlsx':
        raise ReportError('报告文件必须使用 .xlsx 扩展名')
    screenshots = screenshots if screenshots is not None else (task or {}).get('screenshots', [])
    protected = [result.video_metadata.filepath, *protected_paths, *[s['path'] for s in screenshots]]
    for raw in protected:
        source = Path(raw).resolve()
        if source == output or (output.exists() and source.exists() and os.path.samefile(source, output)):
            raise ReportError('报告不能覆盖源视频、数据库或关联截图')
    if output.exists() and not overwrite:
        raise ReportError('目标文件已存在；请选择其他路径或明确允许覆盖')
    check_cancel()
    config = config or (task or {}).get('config') or AnalysisConfig()
    book = Workbook()
    book.remove(book.active)
    temporary = None
    try:
        m = result.video_metadata
        overview = _table(book, '分析概况', ['项目', '内容'], [28, 100])
        fields = [
            ('报告', '智能会议设备投屏质量分析系统 · V0.7'),
            ('导出时间（UTC）', datetime.now(timezone.utc).isoformat(timespec='seconds')),
            ('任务 ID', (task or {}).get('id', '未入库')),
            ('任务名称', (task or {}).get('title', m.filename)),
            ('所属项目', (task or {}).get('project_name') or '未分组'),
            ('任务保存时间（UTC）', (task or {}).get('created_at', '未入库')),
            ('黑屏数', result.statistics['black_count']),
            ('疑似冻结数', result.statistics['freeze_count']),
            ('异常总数', result.statistics['total_count']),
            ('关联截图数', len(screenshots)),
            ('视频名称', m.filename), ('源视频路径', m.filepath),
            ('视频时长（秒）', m.duration), ('分辨率', m.resolution), ('帧率（FPS）', m.fps),
            ('Metadata 帧数', m.frame_count), ('实际处理帧数', result.processed_frames),
            ('实际分析时长（秒）', result.analyzed_duration), ('处理耗时（秒）', result.processing_time),
            ('文件大小（字节）', m.file_size), ('视频编码', m.codec), ('码率（bps）', m.bitrate),
            ('音频编码', m.audio_codec or '无'), ('采样率（Hz）', m.sample_rate or 0),
            ('亮度阈值', config.black_threshold), ('黑屏最短持续（秒）', config.black_duration),
            ('冻结像素差阈值', config.freeze_pixel_threshold), ('冻结变化比例阈值', config.freeze_change_ratio),
            ('冻结最短持续（秒）', config.freeze_duration),
            ('判定说明', '疑似冻结可能由静态 PPT 或低运动内容触发，结果需人工复核。'),
            ('报告性质', '静态结果快照；修改单元格不会重新执行检测或更新汇总。'),
            ('定位与截图', '按固定帧率定位。截图不嵌入报告；本地链接依赖原文件路径，分享时请同时提供截图。'),
        ]
        for row, (label, value) in enumerate(fields, 2):
            _put(overview, row, 1, label)
            _put(overview, row, 2, value)
            overview.row_dimensions[row].height = 34 if isinstance(value, str) and len(value) > 50 else 24
        from core.quality_score import RULE_DESCRIPTION
        q = result.quality
        scoring = [('辅助质量评分（满分100）', q.score if q.score is not None else '不可用'),
                   ('评分规则版本', q.rule_version), ('评分说明', RULE_DESCRIPTION),
                   ('黑屏并集时长（秒）', q.black_seconds),
                   ('非黑屏疑似冻结时长（秒）', q.freeze_seconds),
                   ('黑屏扣分', q.black_penalty), ('疑似冻结扣分', q.freeze_penalty),
                   ('评分边界', q.reason or '项目自定义指标，非设备合格结论；历史结果按当前规则计算。')]
        for row, (label, value) in enumerate(scoring, len(fields)+2):
            _put(overview, row, 1, label)
            _put(overview, row, 2, value)
            overview.row_dimensions[row].height = 48 if label == '评分说明' else 34
        issues = _table(book, '异常明细', ['序号', '类型', '开始时间', '结束时间', '开始（秒）', '结束（秒）', '持续（秒）', '等级', '说明', '检测数据'], [8,18,19,19,14,14,14,12,56,48])
        for row, issue in enumerate(result.issues, 2):
            check_cancel()
            values = [row-1, issue.issue_type, seconds_to_time(issue.start_time), seconds_to_time(issue.end_time),
                      issue.start_time, issue.end_time, issue.duration, issue.severity, issue.description,
                      json.dumps(issue.extra_data or {}, ensure_ascii=False, allow_nan=False)]
            for col, value in enumerate(values, 1):
                _put(issues, row, col, value)
            issues.row_dimensions[row].height = 45
            color = {'HIGH':'FDE8E7','MEDIUM':'FFF3CD','LOW':'EAF2FF'}.get(issue.severity, 'FFFFFF')
            issues.cell(row, 8).fill = PatternFill('solid', fgColor=color)
        if not result.issues:
            _put(issues, 2, 2, '未检测到异常')
        shots = _table(book, '截图关联', ['序号', '异常序号', '帧索引（从0开始）', '时间（秒）', '文件状态', '本地文件路径', '保存时间（UTC）', 'SHA-256'], [8,12,22,14,20,90,30,70])
        for row, shot in enumerate(screenshots, 2):
            check_cancel()
            path = Path(shot['path']).resolve()
            try:
                hasher = hashlib.sha256()
                with path.open('rb') as stream:
                    while chunk := stream.read(1024 * 1024):
                        check_cancel()
                        hasher.update(chunk)
                state = '可用' if hasher.hexdigest() == shot['sha256'] else '内容已更改'
            except OSError:
                state = '缺失或无法读取'
            values = [row-1, shot['issue_ordinal']+1, shot['frame_index'], shot['timestamp'], state,
                      str(path), shot['created_at'], shot['sha256']]
            for col, value in enumerate(values, 1):
                _put(shots, row, col, value)
            shots.row_dimensions[row].height = 45
            if state == '可用':
                shots.cell(row, 6).hyperlink = path.as_uri()
                shots.cell(row, 6).font = Font(name='Microsoft YaHei', size=11, color='245FCC', underline='single')
        if not screenshots:
            _put(shots, 2, 5, '无已关联截图')
        for ws in book:
            ws.auto_filter.ref = ws.dimensions
            ws.print_area = ws.dimensions
        check_cancel()
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(prefix='.report-', suffix='.xlsx', dir=output.parent, delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        check_cancel()
        if overwrite:
            os.replace(temporary, output)
        elif os.name == 'nt':
            # Windows rename 拒绝覆盖已有文件，也支持不提供硬链接的磁盘。
            os.rename(temporary, output)
        else:
            # 原子创建，不因并发创建同名文件而意外覆盖。
            os.link(temporary, output)
        return output
    except ReportError:
        raise
    except Exception as exc:
        raise ReportError(f'报告导出失败：{exc}') from exc
    finally:
        book.close()
        if temporary is not None and temporary.exists():
            temporary.unlink()
